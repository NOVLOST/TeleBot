from aiogram import F, Router,Bot
from aiogram.fsm.context import FSMContext
from aiogram.types import Message, CallbackQuery,ReplyKeyboardRemove
from aiogram.filters import CommandStart, Command
from openpyxl import load_workbook
import openpyxl


import app.keyboards as kb
import classes.client_class as cl
from app.keyboards import start_keyboard

router = Router()
bot = Bot(token='7962043379:AAGXTLRJIlnnDG0nfKHbrGmCkQ_FWo8zdYQ')
redux_row = 0
choice_item = 0
book_name = 'event.xlsx' #название excel файла
sheet_name = 'event' #название листа смотри в самом файле excel
@router.message(CommandStart())
async def cmd_start(message: Message):
    cl.client.full_name = message.from_user.full_name  # Полное имя пользователя

    await message.answer(f"Здравствуйте, {cl.client.full_name}\nЧем я могу вам помочь?", reply_markup=kb.start_keyboard)


@router.message(F.text == 'FAQ ❔')
async def cmd_help(message: Message):
    await message.answer("Пройди короткую регистрацию для участия в конкурсе!")

@router.message(F.text == "Редактировать заявку ✏️")
async def cmd_redux(message: Message,state: FSMContext):
    global redux_row
    book = openpyxl.open('event.xlsx' ,read_only=True)
    sheet = book.active
    flag = False
    for row in range(1,sheet.max_row+1):
        print(sheet[row][9].value)
        if sheet[row][9].value == str(message.from_user.id):
            await message.answer(
                                 f"имя: {sheet[row][0].value}\n"
                                 f"фамилия: {sheet[row][1].value}\n"
                                 f"отчество: {sheet[row][2].value}\n"
                                 f"возраст: {sheet[row][3].value}\n"
                                 f"статус: {sheet[row][4].value}\n"
                                 f"город: {sheet[row][5].value}\n"
                                 f"место обучения: {sheet[row][6].value}\n"
                                 f"номер телефона: {sheet[row][7].value}\n"
                                 f"тип работы: {sheet[row][8].value}"
                                 )
            flag = True
            redux_row = row
            await state.set_state(cl.client.redux_mod)
            await message.answer("Выберите один из пунктов на клавиатуре",reply_markup=kb.redux_bid)
        else:
            continue
    if flag == False:
        await message.answer("У вас нет заявок!")

@router.message(cl.client.redux_mod)
async def redux_mod(message: Message,state: FSMContext):
    global choice_item
    choice_item = int(message.text[0]) - 1
    await state.set_state(cl.client.redux_mod_2)
    await message.answer("Введите новое значение",reply_markup=ReplyKeyboardRemove())

@router.message(cl.client.redux_mod_2)
async def redux_mod(message: Message,state: FSMContext):
    new_value = message.text

    book = load_workbook(book_name)
    sheet = book[sheet_name]
    sheet[redux_row][choice_item].value = new_value
    book.save(book_name)
    book.close()


    await state.clear()
    await message.answer("Успешно! всегда рад помочь вам ^_^",reply_markup=start_keyboard)




#-------------------------------------
#РЕГИСТРАЦИЯ
#-------------------------------------

@router.message(F.text == "Создать новую заявку 📝")
async def registration(message: Message, state: FSMContext):
    await state.update_data(id_account = message.from_user.id)
    print(message.from_user.id)
    await state.set_state(cl.client.real_first_name)
    await message.answer('Введите ваше имя',reply_markup=ReplyKeyboardRemove())

@router.message(cl.client.real_first_name)
async def reg_real_first_name(message: Message, state: FSMContext):
    await state.update_data(real_first_name = message.text )
    data = await state.get_data()
    await state.set_state(cl.client.real_second_name)
    await message.answer("Введите вашу фамилию")

@router.message(cl.client.real_second_name)
async def reg_real_second_name(message: Message, state: FSMContext):
    await state.update_data(real_second_name=message.text)
    await state.set_state(cl.client.real_third_name)
    await message.answer("Введите ваше отчество")

@router.message(cl.client.real_third_name)
async def reg_real_third_name(message: Message, state: FSMContext):
    await state.update_data(real_third_name=message.text)
    await state.set_state(cl.client.school_or_student)
    await message.answer("Вы школьник или студент?",reply_markup=kb.school_of_student)

@router.message(F.text == 'Студент')
async def reg_student(message: Message, state: FSMContext):
    await state.update_data(school_or_student = 'студент')
    await state.set_state(cl.client.city)
    await message.answer("Введите ваш город проживания",reply_markup=ReplyKeyboardRemove())

@router.message(F.text == 'Школьник')
async def reg_school(message: Message, state: FSMContext):
    await state.update_data(school_or_student = 'школьник')
    await state.set_state(cl.client.city)
    await message.answer("Введите ваш город проживания",reply_markup=ReplyKeyboardRemove())

@router.message(cl.client.city)
async def reg_city(message: Message, state: FSMContext):
    await state.update_data(city = message.text)
    await state.set_state(cl.client.education_place)
    await message.answer("Введите ваше место обучение")

@router.message(cl.client.education_place)
async def reg_education_place(message: Message, state: FSMContext):
    await state.update_data(education_place = message.text)
    await state.set_state(cl.client.phone_number)
    await message.answer("Отправьте ваш номер телефона",reply_markup=kb.get_number)

@router.message(cl.client.phone_number,F.contact)
async def reg_number(message: Message, state: FSMContext):
    await state.update_data(phone_number = message.contact.phone_number)
    await state.set_state(cl.client.age)
    await message.answer("Введите ваш возраст")

@router.message(cl.client.phone_number,F.text)
async def reg_number(message: Message, state: FSMContext):
    if len(message.text) == 11:
        await state.update_data(phone_number = message.text)
        await state.set_state(cl.client.age)
        await message.answer("Введите ваш возраст",reply_markup=ReplyKeyboardRemove())
    else:
        await message.answer("Неправильно набран номер")

@router.message(cl.client.age)
async def reg_age(message: Message, state: FSMContext):
    try:
        if int(message.text) < 30:
            await state.update_data(age = message.text)
            await state.set_state(cl.client.type_of_work)
            await message.answer("Выберите тип вашей работы", reply_markup=kb.choice_type_of_work)
        else:
            await message.answer("Введите верный возраст!")
    except:
        await message.answer("Введите верный возраст!")

@router.message(cl.client.type_of_work)
async def reg_work(message: Message, state: FSMContext):
    await state.update_data(type_of_work = message.text)
    await state.set_state(cl.client.id_file)
    await message.answer("Отправте вашу работу")

@router.message(F.photo)
async def photo_handler(message : Message,state: FSMContext):
    photo = message.photo[-1]
    photo_id = photo.file_id
    file_info = await bot.get_file(photo_id)
    print("ITS FILE INFO!!!", file_info.file_path.split('.')[-1])
    format = file_info.file_path.split('.')[-1]
    data = await state.get_data()
    cl.client.num_of_photo += 1
    await bot.download_file(file_info.file_path,f'{cl.client.num_of_photo}_{data["real_first_name"]} {data["real_second_name"]} {data["id_account"]}.{format}')
    await message.answer(f"имя: {data['real_first_name']}\n"
                         f" фамилия: {data['real_second_name']}\n"
                         f" отчество: {data['real_third_name']}\n"
                         f" возраст: {data['age']}\n"
                         f" статус: {data['school_or_student']}\n"
                         f" город: {data['city']}\n"
                         f" место обучения: {data['education_place']}\n"
                         f" телефон: {data['phone_number']}\n"
                         f" тип работы: {data['type_of_work']}")



    wb = load_workbook(book_name)
    ws = wb[sheet_name]

    ws.append([f" {data['real_first_name']}",
                         f" {data['real_second_name']}",
                         f" {data['real_third_name']}",
                         f" {data['age']}",
                         f" {data['school_or_student']}",
                         f" {data['city']}",
                         f" {data['education_place']}",
                         f" {data['phone_number']}",
                         f" {data['type_of_work']}",
                         f" {data['id_account']}"[1:]
               ])
    wb.save(book_name)
    wb.close()



    await message.answer("Все верно? если допустили ошибку попроси меня для редактирования заявки ;)", reply_markup=kb.start_keyboard)
    await state.clear()


@router.message(F.document)
async def photo_handler(message : Message,state: FSMContext):
    doc = message.document
    doc_id = doc.file_id
    file_info = await bot.get_file(doc_id)
    print("ITS FILE INFO!!!",file_info.file_path.split('.'))
    format = file_info.file_path.split('.')[-1]
    data = await state.get_data()
    cl.client.num_of_photo += 1
    await bot.download_file(file_info.file_path,f'{cl.client.num_of_photo}_{data["real_first_name"]} {data["real_second_name"]} {data["id_account"]}.{format}')
    await message.answer(f"имя: {data['real_first_name']}\n"
                         f" фамилия: {data['real_second_name']}\n"
                         f" отчество: {data['real_third_name']}\n"
                         f" возраст: {data['age']}\n"
                         f" статус: {data['school_or_student']}\n"
                         f" город: {data['city']}\n"
                         f" место обучения: {data['education_place']}\n"
                         f" телефон: {data['phone_number']}\n"
                         f" тип работы: {data['type_of_work']}")



    wb = load_workbook(book_name)
    ws = wb[sheet_name]

    ws.append([f" {data['real_first_name']}",
                         f" {data['real_second_name']}",
                         f" {data['real_third_name']}",
                         f" {data['age']}",
                         f" {data['school_or_student']}",
                         f" {data['city']}",
                         f" {data['education_place']}",
                         f" {data['phone_number']}",
                         f" {data['type_of_work']}",
                         f" {data['id_account']}"[1:]
               ])
    wb.save(book_name)
    wb.close()



    await message.answer("Все верно? если допустили ошибку попроси меня для редактирования заявки ;)", reply_markup=kb.start_keyboard)
    await state.clear()

@router.message(F.audio)
async def photo_handler(message : Message,state: FSMContext):
    doc = message.audio
    doc_id = doc.file_id
    file_info = await bot.get_file(doc_id)
    print("ITS FILE INFO!!!",file_info.file_path.split('.'))
    format = file_info.file_path.split('.')[-1]
    data = await state.get_data()
    cl.client.num_of_photo += 1
    await bot.download_file(file_info.file_path,f'{cl.client.num_of_photo}_{data["real_first_name"]} {data["real_second_name"]} {data["id_account"]}.{format}')
    await message.answer(f"имя: {data['real_first_name']}\n"
                         f" фамилия: {data['real_second_name']}\n"
                         f" отчество: {data['real_third_name']}\n"
                         f" возраст: {data['age']}\n"
                         f" статус: {data['school_or_student']}\n"
                         f" город: {data['city']}\n"
                         f" место обучения: {data['education_place']}\n"
                         f" телефон: {data['phone_number']}\n"
                         f" тип работы: {data['type_of_work']}")



    wb = load_workbook(book_name)
    ws = wb[sheet_name]

    ws.append([f" {data['real_first_name']}",
                         f" {data['real_second_name']}",
                         f" {data['real_third_name']}",
                         f" {data['age']}",
                         f" {data['school_or_student']}",
                         f" {data['city']}",
                         f" {data['education_place']}",
                         f" {data['phone_number']}",
                         f" {data['type_of_work']}",
                         f" {data['id_account']}"[1:]
               ])
    wb.save(book_name)
    wb.close()



    await message.answer("Все верно? если допустили ошибку попроси меня для редактирования заявки ;)", reply_markup=kb.start_keyboard)
    await state.clear()

@router.message(F.video)
async def photo_handler(message : Message,state: FSMContext):
    doc = message.video
    doc_id = doc.file_id
    file_info = await bot.get_file(doc_id)
    print("ITS FILE INFO!!!",file_info.file_path.split('.'))
    format = file_info.file_path.split('.')[-1]
    data = await state.get_data()
    cl.client.num_of_photo += 1
    await bot.download_file(file_info.file_path,f'{cl.client.num_of_photo}_{data["real_first_name"]} {data["real_second_name"]} {data["id_account"]}.{format}')
    await message.answer(f"имя: {data['real_first_name']}\n"
                         f" фамилия: {data['real_second_name']}\n"
                         f" отчество: {data['real_third_name']}\n"
                         f" возраст: {data['age']}\n"
                         f" статус: {data['school_or_student']}\n"
                         f" город: {data['city']}\n"
                         f" место обучения: {data['education_place']}\n"
                         f" телефон: {data['phone_number']}\n"
                         f" тип работы: {data['type_of_work']}")



    wb = load_workbook(book_name)
    ws = wb[sheet_name]

    ws.append([f" {data['real_first_name']}",
                         f" {data['real_second_name']}",
                         f" {data['real_third_name']}",
                         f" {data['age']}",
                         f" {data['school_or_student']}",
                         f" {data['city']}",
                         f" {data['education_place']}",
                         f" {data['phone_number']}",
                         f" {data['type_of_work']}",
                         f" {data['id_account']}"[1:]
               ])
    wb.save(book_name)
    wb.close()



    await message.answer("Все верно? если допустили ошибку попроси меня для редактирования заявки ;)", reply_markup=kb.start_keyboard)
    await state.clear()
